"""Duplicate REPORT.md as an editable Word document with native equations.

python build_word.py [--math-xsl PATH]
Microsoft Office supplies MML2OMML.XSL for MathML-to-Word-equation conversion.
"""

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import latex2mathml.converter
from lxml import etree, html
import markdown

from artifact_checks import verify_artifacts

ROOT = Path(__file__).resolve().parent
WORD_NAME = "Albert_Kabore_Project1_Report.docx"
BLACK = RGBColor(0, 0, 0)


def find_math_transform(explicit=None):
    candidates = [Path(explicit)] if explicit else [
        Path("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"),
        Path("C:/Program Files (x86)/Microsoft Office/root/Office16/MML2OMML.XSL"),
        Path("C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return etree.XSLT(etree.parse(str(candidate)))
    raise FileNotFoundError("Pass --math-xsl with the path of Office's MML2OMML.XSL")


def format_run(run, bold=False, italic=False, code=False, size=None):
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = BLACK
    run.font.name = "Consolas" if code else "Calibri"
    if size or code:
        run.font.size = Pt(size or 9)
    return run


def add_hyperlink(paragraph, text, target, bold=False, italic=False):
    relation = paragraph.part.relate_to(
        target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    for enabled, tag in [(bold, "w:b"), (italic, "w:i")]:
        if enabled:
            props.append(OxmlElement(tag))
    run.append(props)
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def inline(element, paragraph, bold=False, italic=False, code=False, size=None):
    if element.text:
        format_run(paragraph.add_run(element.text), bold, italic, code, size)
    for child in element:
        tag = child.tag
        if tag == "br":
            paragraph.add_run().add_break()
        elif tag == "a":
            add_hyperlink(paragraph, "".join(child.itertext()), child.get("href"), bold, italic)
        elif tag == "img":
            path = (ROOT / child.get("src")).resolve()
            path.relative_to(ROOT)
            picture = paragraph.add_run().add_picture(str(path), width=Inches(6.8))
            picture._inline.docPr.set("descr", child.get("alt", ""))
        else:
            inline(child, paragraph, bold or tag in {"strong", "b"},
                   italic or tag in {"em", "i"}, code or tag == "code", size)
        if child.tail:
            format_run(paragraph.add_run(child.tail), bold, italic, code, size)


def add_equation(document, expression, transform):
    mathml = latex2mathml.converter.convert(expression, display="block")
    equation = deepcopy(transform(etree.fromstring(mathml.encode("utf-8"))).getroot())
    if equation.tag == qn("m:oMath"):
        wrapper = OxmlElement("m:oMathPara")
        wrapper.append(equation)
        equation = wrapper
    # Direct black formatting also applies to native mathematical runs.
    for run in equation.iter(qn("m:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(1 if run.find(qn("m:rPr")) is not None else 0, properties)
        color = properties.find(qn("w:color"))
        if color is None:
            color = OxmlElement("w:color")
            properties.append(color)
        color.set(qn("w:val"), "000000")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph._p.append(equation)


def add_table(document, element):
    rows = element.findall(".//tr")
    columns = len(rows[0])
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    proportions = [1 / columns] * columns
    if columns == 2:
        proportions = [0.37, 0.63]
    elif columns == 3:
        proportions = [0.28, 0.29, 0.43]
    elif columns == 4:
        proportions = [0.15, 0.35, 0.25, 0.25]
    elif columns == 9:
        proportions = [0.06] + [0.1175] * 8
    for col, proportion in zip(table.columns, proportions):
        col.width = Inches(7.14 * proportion)
    for row_index, source_row in enumerate(rows):
        row = table.add_row()
        tr_properties = row._tr.get_or_add_trPr()
        tr_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            tr_properties.append(OxmlElement("w:tblHeader"))
        for cell, source_cell, proportion in zip(row.cells, source_row, proportions):
            cell.width = Inches(7.14 * proportion)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            inline(source_cell, paragraph, bold=row_index == 0, size=8 if columns >= 7 else 9)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document):
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.68)
    for style in document.styles:
        if style.type in {WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER}:
            style.font.color.rgb = BLACK
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    if "Code Block" not in document.styles:
        style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Consolas"
        style.font.size = Pt(8.5)
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = 1.0
    document.core_properties.title = "Project 1: Logistic Regression and Feed-Forward Neural Network"
    document.core_properties.author = "Albert Kabore"
    document.core_properties.subject = "Sentiment classification with six features"


def build_word(math_xsl=None):
    results = json.loads((ROOT / "results/results.json").read_text(encoding="utf-8"))
    verify_artifacts(results)
    source = (ROOT / "REPORT.md").read_text(encoding="utf-8")
    equations = []

    def replace_equation(match):
        equations.append(match.group(1).strip())
        return f'\n\n<div data-equation="{len(equations)-1}"></div>\n\n'

    prepared = re.sub(r"\$\$(.*?)\$\$", replace_equation, source, flags=re.DOTALL)
    tree = html.fragment_fromstring(markdown.markdown(prepared, extensions=["tables", "fenced_code"]), create_parent="div")
    transform = find_math_transform(math_xsl)
    document = Document()
    configure_document(document)
    for element in tree:
        tag = element.tag
        if tag in {"h1", "h2", "h3"}:
            paragraph = document.add_paragraph(style={"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}[tag])
            inline(element, paragraph, bold=True)
        elif tag == "div" and element.get("data-equation") is not None:
            add_equation(document, equations[int(element.get("data-equation"))], transform)
        elif tag == "table":
            add_table(document, element)
        elif tag == "pre":
            paragraph = document.add_paragraph(style="Code Block")
            format_run(paragraph.add_run("".join(element.itertext()).rstrip()), code=True, size=8.5)
        elif tag in {"ol", "ul"}:
            for index, item in enumerate(element.findall("li"), start=1):
                paragraph = document.add_paragraph()
                format_run(paragraph.add_run(f"{index}. " if tag == "ol" else "• "))
                inline(item, paragraph)
        elif tag == "blockquote":
            for quote in element:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.15)
                inline(quote, paragraph, italic=True)
        elif tag == "p":
            paragraph = document.add_paragraph()
            inline(element, paragraph)
            if element.find("img") is not None:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            raise ValueError(f"Unsupported report element: {tag}")
    destination = ROOT / WORD_NAME
    document.save(destination)
    print(f"Created {destination.name}: {len(equations)} editable equations, {len(document.tables)} tables and {len(document.inline_shapes)} figures.")
    return destination


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--math-xsl", help="Path to Microsoft's MML2OMML.XSL")
    build_word(parser.parse_args().math_xsl)
